# -*- coding: utf-8 -*-
"""实验报告 DOCX 的排版工具（从 build_report.py 抽出，供 build_report3.py 复用）。

版式要求：宋体、小四（12pt）、1.25 倍行距、A4。
反引号 ` 包裹的片段用等宽字体排版，代码块用灰色底纹的等宽字体。
"""
import re

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "宋体"
CODE_FONT = "Consolas"
BODY_SIZE = 12          # 小四
CODE_SIZE = 9


def set_run_font(run, name=BODY_FONT, size=BODY_SIZE, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def shade(paragraph, fill="F2F2F2"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def fmt_para(p, indent=True, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.25,
             before=0, after=0):
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    p.alignment = align
    pf.first_line_indent = Pt(BODY_SIZE * 2) if indent else Pt(0)
    pf.widow_control = True
    return p


def add_rich(cell, text, indent=True, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
             spacing=1.25, before=0, after=0, size=BODY_SIZE):
    """写入一段文字，反引号内的内容用等宽字体。"""
    p = cell.add_paragraph()
    fmt_para(p, indent=indent, align=align, spacing=spacing, before=before, after=after)
    for i, seg in enumerate(text.split("`")):
        if not seg:
            continue
        if i % 2 == 1:
            r = p.add_run(seg)
            set_run_font(r, CODE_FONT, size - 2, bold)
            r.font.color.rgb = RGBColor(0xA3, 0x1F, 0x1F)
        else:
            r = p.add_run(seg)
            set_run_font(r, BODY_FONT, size, bold)
    return p


def add_code(cell, code, size=CODE_SIZE):
    """等宽字体的代码/输出块，灰色底纹。"""
    lines = code.strip("\n").split("\n")
    p = cell.add_paragraph()
    fmt_para(p, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0, before=2, after=2)
    p.paragraph_format.left_indent = Cm(0.35)
    shade(p)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, CODE_FONT, size)
    return p


def add_heading(cell, text, size=BODY_SIZE, before=6, after=2):
    p = cell.add_paragraph()
    fmt_para(p, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.25,
             before=before, after=after)
    r = p.add_run(text)
    set_run_font(r, BODY_FONT, size, bold=True)
    return p


def set_col_width(table, widths_cm):
    """显式写入 tblW / tblGrid / 每个 tc 的 tcW，保证 Word 与 LibreOffice 一致按固定布局排版。"""
    tw = lambda cm: int(round(cm / 2.54 * 1440))   # cm -> twips
    total = tw(sum(widths_cm))

    table.autofit = False
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(el)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.insert(0, layout)

    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.insert(0, tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(total))
    for el in tblPr.findall(qn("w:tblInd")):
        tblPr.remove(el)
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:type"), "dxa")
    ind.set(qn("w:w"), "0")
    tblPr.append(ind)

    grid = table._tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(tw(w)))

    for row in table.rows:
        col = 0
        for tc in row._tr.findall(qn("w:tc")):
            tcPr = tc.get_or_add_tcPr()
            span = tcPr.find(qn("w:gridSpan"))
            n = int(span.get(qn("w:val"))) if span is not None else 1
            w = tw(sum(widths_cm[col:col + n]))
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.insert(0, tcW)
            tcW.set(qn("w:type"), "dxa")
            tcW.set(qn("w:w"), str(w))
            col += n


def set_borders(table, sz=6):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("第 ")
    set_run_font(r, BODY_FONT, 9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rPr.append(sz)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    p._p.append(fld)
    r2 = p.add_run(" 页")
    set_run_font(r2, BODY_FONT, 9)


def put(cell, text, **kw):
    """首段复用 cell 的第 0 段，之后追加新段。"""
    if cell.paragraphs[0].text == "" and not cell.paragraphs[0].runs:
        p = cell.paragraphs[0]
        fmt_para(p, indent=kw.get("indent", True),
                 align=kw.get("align", WD_ALIGN_PARAGRAPH.LEFT),
                 spacing=kw.get("spacing", 1.25), before=kw.get("before", 0),
                 after=kw.get("after", 0))
        for i, seg in enumerate(text.split("`")):
            if not seg:
                continue
            if i % 2 == 1:
                r = p.add_run(seg)
                set_run_font(r, CODE_FONT, kw.get("size", BODY_SIZE) - 2, kw.get("bold", False))
                r.font.color.rgb = RGBColor(0xA3, 0x1F, 0x1F)
            else:
                r = p.add_run(seg)
                set_run_font(r, BODY_FONT, kw.get("size", BODY_SIZE), kw.get("bold", False))
        return p
    return add_rich(cell, text, **kw)


def put_code(cell, code, size=CODE_SIZE):
    if cell.paragraphs[0].text == "" and not cell.paragraphs[0].runs:
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    return add_code(cell, code, size)


def put_heading(cell, text, **kw):
    if cell.paragraphs[0].text == "" and not cell.paragraphs[0].runs:
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    return add_heading(cell, text, **kw)


def iter_all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def apply_keep_rules(doc):
    """标题与其后正文同页；代码/输出块不跨页断开。"""
    pat = re.compile(r"^(一、|二、|三、|四、|五、|问题[一二三四])")
    for p in iter_all_paragraphs(doc):
        text = p.text.strip()
        pf = p.paragraph_format
        if pat.match(text):
            pf.keep_with_next = True
            pf.keep_together = True
        elif p.runs and all((r.font.name or "") == CODE_FONT for r in p.runs if r.text):
            pf.keep_together = True
            pf.keep_with_next = True


def new_report_doc():
    """新建 A4 页面、Normal 样式为宋体小四、页脚带页码的文档。"""
    from docx import Document
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    add_page_number_footer(sec)
    return doc


def add_title(doc, text="计算机科学与技术学院  大数据管理与分析  课程实验报告"):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    r = title.add_run(text)
    set_run_font(r, BODY_FONT, 15, bold=True)
    return title


def new_report_table(doc, widths=(4.5, 5.75, 5.75)):
    from docx.enum.table import WD_ALIGN_VERTICAL
    table = doc.add_table(rows=9, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(table)
    set_col_width(table, list(widths))
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return table


def add_tail_paragraph(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("")
    set_run_font(r, BODY_FONT, 6)
    return p
